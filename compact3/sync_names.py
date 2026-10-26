#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-License-Identifier: GPL-3.0-only
"""Rewrite story speaker prefixes from the 名称_* sheets.
Build-time apply_name_rules() does the same to the ROM; this keeps the
workbook display in sync after a name is edited."""
import re
from pathlib import Path
import openpyxl

XLSX = Path(__file__).resolve().parent / "assets" / "text.xlsx"
SPEAKER_RE = re.compile(r"^([^「（(]{1,16})([「（(])")
SKIP = {"強化パーツ", "ユニット", "特殊技能"}

wb = openpyxl.load_workbook(XLSX)
names = []
for sn in wb.sheetnames:
    if not sn.startswith("名称_"):
        continue
    ws = wb[sn]
    hdr = [c.value for c in ws[1]]
    if "index" not in hdr:
        continue
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(hdr, row))
        if r.get("index") is None or not str(r.get("zh") or "").strip():
            continue
        names.append((int(r["index"]), str(r.get("jp") or "").strip(), str(r["zh"]).strip(),
                      str(r.get("别名") or "").strip()))
names.sort()
jp_map = {}
for i, j, z, a in names:
    if i >= 331 or j not in jp_map:
        jp_map[j] = z
alias_pairs = []
for i, j, z, a in names:
    if not z or not a:
        continue
    for old in re.split(r"[\s,，、;；]+", a):
        old = old.strip()
        if len(old) >= 2 and old != z:
            alias_pairs.append((old, z))
alias_pairs.sort(key=lambda p: -len(p[0]))

changed = alias_hits = 0
for sn in wb.sheetnames:
    if not sn.startswith("剧情_"):
        continue
    ws = wb[sn]
    hdr = [c.value for c in ws[1]]
    ji = hdr.index("jp"); zi = hdr.index("zh"); ei = hdr.index("empty")
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, ei + 1).value or "") == "1":
            continue
        j = str(ws.cell(r, ji + 1).value or "")
        z = str(ws.cell(r, zi + 1).value or "")
        for old, cur in alias_pairs:
            if old in z:
                alias_hits += z.count(old)
                z = z.replace(old, cur)
        mj, mz = SPEAKER_RE.match(j), SPEAKER_RE.match(z)
        if not mj or not mz:
            continue
        sp = mj.group(1).strip()
        if sp in SKIP or sp not in jp_map or mz.group(1) == jp_map[sp]:
            continue
        z = jp_map[sp] + mz.group(2) + z[mz.end():]
        changed += 1
        ws.cell(r, zi + 1).value = z
    # write alias-only rows too
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, ei + 1).value or "") == "1":
            continue
        cur = str(ws.cell(r, zi + 1).value or "")
        new = cur
        for old, val in alias_pairs:
            if old in new:
                new = new.replace(old, val)
        if new != cur:
            ws.cell(r, zi + 1).value = new
wb.save(XLSX)
print("speaker names synced:", changed, "alias hits:", alias_hits)
