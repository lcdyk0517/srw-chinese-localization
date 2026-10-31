#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-License-Identifier: GPL-3.0-only
"""text.xlsx 与 assets/text/*.tsv 纯文本镜像的双向同步。

git 只保存 TSV（每个 sheet 一个文件，diff 逐行可见）；text.xlsx 移出版本库。
build.py 构建时自动从 TSV 重建 assets/text.build.xlsx，原有解析逻辑不变。

用法：
  python text_sync.py export          # assets/text.xlsx → assets/text/*.tsv（在 Excel 里改完后运行，然后提交 TSV）
  python text_sync.py xlsx [--force]  # assets/text/*.tsv → 重建 assets/text.xlsx（新 clone 后生成可编辑工作簿）
  python text_sync.py check           # 逐格比对 assets/text.xlsx 与 TSV，不一致时退出码 1

TSV 格式：UTF-8、制表符分隔、首行为表头；单元格内的反斜杠、Tab、CR、LF
转义为 \\\\、\\t、\\r、\\n；空字段即空单元格（重建时写回 None）。
数值列（数据区全部为数字的列）记录在 _manifest.json，重建时回写为数字类型；
混排列（数字 + 文本）一律按文本重建，build.py 的解析对两者均兼容。
"""
import argparse
import json
import os
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
XLSX = HERE / "assets" / "text.xlsx"
BUILD_XLSX = HERE / "assets" / "text.build.xlsx"
TSV_DIR = HERE / "assets" / "text"
MANIFEST = TSV_DIR / "_manifest.json"
MTIME_SLACK = 2  # 秒；mtime 比较的容差


def _esc(s: str) -> str:
    return (s.replace("\\", "\\\\").replace("\t", "\\t")
             .replace("\r", "\\r").replace("\n", "\\n"))


def _unesc(s: str) -> str:
    out, i = [], 0
    while i < len(s):
        c = s[i]
        if c == "\\" and i + 1 < len(s) and s[i + 1] in "trn\\":
            out.append({"t": "\t", "r": "\r", "n": "\n", "\\": "\\"}[s[i + 1]])
            i += 2
        else:
            out.append(c)
            i += 1
    return "".join(out)


def _cell_text(v) -> str:
    if v is None or v == "":
        return ""
    return _esc(str(v))


def _read_grid(ws):
    """整表读为二维网格：去掉尾部全空行与尾部全空列。"""
    grid = [tuple(row) for row in ws.iter_rows(values_only=True)]
    while grid and all(v is None for v in grid[-1]):
        grid.pop()
    width = 0
    for row in grid:
        for i in range(len(row) - 1, -1, -1):
            if row[i] is not None:
                width = max(width, i + 1)
                break
    return [row[:width] for row in grid]


def _scan_formulas(path: Path):
    """返回 (公式总数, 缓存值非空的公式数)。后者导出会丢值，硬失败。"""
    import re
    total = live = 0
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("xl/worksheets/") and n.endswith(".xml"):
                b = z.read(n)
                cells = re.findall(rb"<c [^>]*>.*?</c>", b)
                for c in cells:
                    if re.search(rb"<f[ />]", c):
                        total += 1
                        m = re.search(rb"<f[^>]*>.*?</f><v>(.*?)</v>", c, re.S)
                        if m and m.group(1).strip():
                            live += 1
    return total, live


def export(xlsx: Path = XLSX, tsv_dir: Path = TSV_DIR) -> None:
    if not xlsx.is_file():
        sys.exit(f"text workbook not found: {xlsx}")
    n_formula, n_live = _scan_formulas(xlsx)
    if n_live:
        sys.exit(f"{xlsx} 含 {n_live} 个带缓存值的公式单元格，TSV 镜像只保存值，请先移除公式")
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    tsv_dir.mkdir(parents=True, exist_ok=True)
    for p in tsv_dir.glob("*.tsv"):
        p.unlink()  # sheet 改名后清掉旧文件
    sheets, n_emptystr = [], 0
    for name in wb.sheetnames:
        ws = wb[name]
        grid = _read_grid(ws)
        n_emptystr += sum(1 for row in grid for v in row if v == "")
        header = grid[0] if grid else ()
        data = grid[1:]
        numeric = {}
        for ci in range(len(header)):
            kinds = {type(row[ci]).__name__ for row in data
                     if ci < len(row) and row[ci] is not None}
            if kinds == {"int"}:
                numeric[ci + 1] = "int"
            elif kinds and kinds <= {"int", "float"}:
                numeric[ci + 1] = "float"
        body = "\n".join("\t".join(_cell_text(v) for v in row) for row in grid)
        (tsv_dir / f"{name}.tsv").write_text(
            body + "\n" if body else "", encoding="utf-8", newline="\n")
        sheets.append({
            "name": name,
            "numeric_cols": [[ci, kind] for ci, kind in sorted(numeric.items())],
            "merged": [str(r) for r in ws.merged_cells.ranges],
        })
    (tsv_dir / "_manifest.json").write_text(
        json.dumps({"version": 1, "sheets": sheets}, ensure_ascii=False, indent=1) + "\n",
        encoding="utf-8", newline="\n")
    if n_emptystr:
        print(f"[警告] 有 {n_emptystr} 个空字符串单元格，已按空单元格导出（重建后变为 None）")
    if n_formula:
        print(f"[提示] {n_formula} 个空缓存公式单元格按空值导出（重建后变为普通空单元格，与构建行为一致）")
    print(f"exported {len(sheets)} sheets -> {tsv_dir}")


def write_xlsx(tsv_dir: Path = TSV_DIR, out: Path = BUILD_XLSX) -> Path:
    """从 TSV 重建工作簿；mtime 回拨到最新 TSV，避免触发“未导出”误报。"""
    manifest = json.loads((tsv_dir / "_manifest.json").read_text(encoding="utf-8"))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for entry in manifest["sheets"]:
        name = entry["name"]
        numeric = {int(ci): kind for ci, kind in entry.get("numeric_cols", [])}
        ws = wb.create_sheet(title=name)
        lines = (tsv_dir / f"{name}.tsv").read_text(encoding="utf-8").split("\n")
        while lines and lines[-1] == "":
            lines.pop()
        lines = [ln[:-1] if ln.endswith("\r") else ln for ln in lines]  # 兼容被检出为 CRLF 的情况
        for ri, line in enumerate(lines):
            row = []
            for ci, field in enumerate(line.split("\t")):
                v = _unesc(field) if field else None
                if v is not None and ri > 0 and (ci + 1) in numeric:
                    v = int(v) if numeric[ci + 1] == "int" else float(v)
                row.append(v)
            ws.append(row)
        for rng in entry.get("merged", []):
            ws.merge_cells(rng)
    wb.save(out)
    newest = max((p.stat().st_mtime for p in tsv_dir.glob("*.tsv")), default=None)
    if newest is not None:
        os.utime(out, (newest, newest))
    return out


def stale_hint() -> None:
    """assets/text.xlsx 比 TSV 新时提醒 export（xlsx 是 gitignored 的编辑面）。"""
    if not (TSV_DIR.is_dir() and XLSX.is_file()):
        return
    newest = max((p.stat().st_mtime for p in TSV_DIR.glob("*.tsv")), default=0)
    if XLSX.stat().st_mtime > newest + MTIME_SLACK:
        print("[提示] assets/text.xlsx 比 assets/text/*.tsv 新：如在 Excel 里改过，"
              "请运行 python text_sync.py export 同步后再提交")


def check() -> int:
    """把本地 xlsx 导出到临时目录，与入库 TSV 逐字节比对。"""
    tmp = HERE / "_tsv_check_tmp"
    export(XLSX, tmp)
    names = ({p.name for p in tmp.glob("*.tsv")} | {p.name for p in TSV_DIR.glob("*.tsv")}
             | {"_manifest.json"})
    diffs = []
    for n in sorted(names):
        a, b = tmp / n, TSV_DIR / n
        if not a.is_file() or not b.is_file():
            diffs.append(f"{n}: 仅存在于 {'xlsx 导出' if a.is_file() else 'TSV 目录'}")
            continue
        ta = a.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
        tb = b.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
        if ta != tb:
            for i, (x, y) in enumerate(zip(ta, tb), 1):
                if x != y:
                    diffs.append(f"{n}:{i}\n  xlsx: {x[:120]}\n  tsv : {y[:120]}")
                    break
            else:
                diffs.append(f"{n}: 行数不同 ({len(ta)} vs {len(tb)})")
    shutil.rmtree(tmp, ignore_errors=True)
    if diffs:
        print("TSV 与 assets/text.xlsx 不一致（在 Excel 里改过？先运行 python text_sync.py export）")
        for d in diffs[:20]:
            print(" ", d)
        return 1
    print("TSV 与 assets/text.xlsx 一致")
    return 0


def cmd_xlsx(force: bool) -> int:
    if XLSX.is_file() and not force:
        newest = max((p.stat().st_mtime for p in TSV_DIR.glob("*.tsv")), default=0)
        if XLSX.stat().st_mtime > newest + MTIME_SLACK:
            print("assets/text.xlsx 比 TSV 新，直接覆盖会丢失未导出的修改；"
                  "先运行 python text_sync.py export，或加 --force")
            return 1
    print("written:", write_xlsx(TSV_DIR, XLSX))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="text.xlsx 与 TSV 镜像同步")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("export", help="assets/text.xlsx -> assets/text/*.tsv")
    p_x = sub.add_parser("xlsx", help="assets/text/*.tsv -> 重建 assets/text.xlsx")
    p_x.add_argument("--force", action="store_true", help="即使 xlsx 比 TSV 新也覆盖")
    sub.add_parser("check", help="比对 xlsx 与 TSV 是否一致")
    args = ap.parse_args()
    if args.cmd == "export":
        export()
        return 0
    if args.cmd == "xlsx":
        return cmd_xlsx(args.force)
    return check()


if __name__ == "__main__":
    sys.exit(main())
